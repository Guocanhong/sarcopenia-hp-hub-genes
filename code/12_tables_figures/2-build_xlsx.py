# -*- coding: utf-8 -*-
"""Build the four manuscript tables into one submission-ready xlsx workbook."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"D:\GEO-limma,肌少症1\17-Tables"
OUT = BASE + r"\四张表格_投稿版.xlsx"

wb = Workbook()
wb.remove(wb.active)

TITLE_FONT = Font(name="Arial", size=11, bold=True)
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="555555")
HEAD_FILL = PatternFill("solid", fgColor="2F5496")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def fmt(v):
    if v is None or v == "" or v == "NA":
        return "—"
    if isinstance(v, str):
        try:
            v = float(v)
        except ValueError:
            return v
    if isinstance(v, float):
        if v != 0 and (abs(v) < 1e-3 or abs(v) >= 1e5):
            return "%.2e" % v
        return round(v, 3)
    return v


def write_sheet(ws, title, header, rows, notes, col_widths=None):
    ws.append([title])
    ws["A1"].font = TITLE_FONT
    ws.append([])
    ws.append(header)
    hr = ws.max_row
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for r in rows:
        ws.append([fmt(v) for v in r])
        rr = ws.max_row
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            if isinstance(cell.value, str) and len(cell.value) > 60:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    for n in notes:
        ws.append([n])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT
    widths = col_widths or [max(len(str(header[i])) for i in range(len(header))) + 4]
    # auto width from content
    for c in range(1, len(header) + 1):
        w = len(str(header[c - 1]))
        for r in range(hr + 1, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                w = max(w, min(len(str(v)), 55))
        ws.column_dimensions[get_column_letter(c)].width = min(w + 3, 58)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)


# ---------------- Sheet 1: Table 1 ----------------
ws = wb.create_sheet("Table1_Datasets")
header = ["Role", "Dataset", "Tissue", "Platform", "Groups (n)", "Purpose in this study"]
rows = [
    ["Training / DEG (sarcopenia)", "GSE1428", "Vastus lateralis, skeletal muscle",
     "GPL96 (Affymetrix)", "Young 10 vs Old 12",
     "DEG screening, WGCNA, ML feature selection, ROC (training)"],
    ["Training / DEG (H. pylori)", "GSE27411",
     "Gastric mucosa (HP-negative vs HP-positive non-atrophic; 3 donors/group, antrum + corpus)",
     "GPL6255 (Illumina humanRef-8)", "Control 6 vs HP 6 biopsies",
     "DEG screening, WGCNA, ML feature selection, ROC (training)"],
    ["External validation (sarcopenia, RNA-seq)", "GSE111016 (Singapore Sarcopenia Study)",
     "Vastus lateralis, skeletal muscle", "RNA-seq", "Control 20 vs SP 19 (1 SP sample excluded: all genes missing in processed matrix)",
     "External ROC validation, consensus clustering"],
    ["External validation (H. pylori)", "GSE5081", "Gastric mucosa",
     "GPL570 (Affymetrix HG-U133 Plus 2.0)", "Control 16 vs HP 16",
     "External ROC validation"],
    ["External validation (ageing)", "GSE25941", "Skeletal muscle",
     "GPL570 (Affymetrix HG-U133 Plus 2.0)", "Young 15 vs Old 21",
     "External ROC validation (ageing design)"],
    ["External validation (sarcopenia, bulk RNA-seq)", "GSE167186 series (GPL20301 platform)",
     "Skeletal muscle (vastus lateralis)", "Bulk RNA-seq",
     "Sarcopenia 24 / Old 29 / Young 19 (4 UNCLASSIFIED excluded)",
     "Direction replication, ROC, phenotype correlation"],
    ["Single-cell (skeletal muscle)",
     "Human Skeletal Muscle Ageing Cell Atlas (Kedlian et al., Nat Aging 2024)",
     "Skeletal muscle", "10x scRNA/snRNA",
     "Young 8 donors (15-40 y; 6M/2F) vs Old 9 donors (50-75 y; 4M/5F); 183,161 cells/nuclei",
     "Cell-type localization of hub genes"],
    ["Single-cell (stomach)", "GSE249874", "Gastric mucosa (non-atrophic gastritis)",
     "10x scRNA", "Hp-negative 3 vs Hp-positive 3 adult donors (age/sex not provided in GEO); 46,753 cells before / 26,250 after QC",
     "Cell-type localization in H. pylori line"],
    ["Protein level", "Human Protein Atlas (HPA)", "Skeletal muscle / stomach",
     "Immunohistochemistry", "Not applicable",
     "Protein-level corroboration of hub genes"],
]
notes = [
    "Note: GSE1428 and GSE25941 are young-vs-old ageing designs rather than AWGS-defined sarcopenia cohorts; "
    "this is stated explicitly in Methods and Limitations.",
    "All datasets are publicly available; no new human samples were collected in this study.",
]
write_sheet(ws, "Table 1. Summary of datasets used in this study.", header, rows, notes)

# ---------------- Sheet 2: Table 2 ----------------
ws = wb.create_sheet("Table2_118DEGs")
data = read_csv(BASE + r"\output\Table2_118DEGs_data.csv")
header = ["Gene", "logFC (SP)", "P (SP)", "adj. P (SP)", "Reg. (SP)",
          "logFC (HP)", "P (HP)", "adj. P (HP)", "Reg. (HP)"]
rows = data[1:]
notes = [
    "SP = sarcopenia line (GSE1428, SP vs Control); HP = H. pylori line (GSE27411, HP vs Control).",
    "Reg. = regulation direction (Up/Down); logFC = log2 fold change.",
    "Genes differentially expressed in both lines (P < 0.05), n = 118.",
]
write_sheet(ws, "Table 2. The 118 genes differentially expressed in both the sarcopenia and H. pylori lines.",
            header, rows, notes)

# ---------------- Sheet 3: Table 3 ----------------
ws = wb.create_sheet("Table3_HubGenes")
data = read_csv(BASE + r"\output\Table3_hubgene_stats_data.csv")
header = ["Dataset", "Comparison", "Gene", "Direction", "AUC", "95% CI low",
          "95% CI high", "P (Wilcoxon)", "n ref", "n test"]
rows = data[1:]
notes = [
    "Direction = expression direction in the test group vs reference group (Up/Down, by group median).",
    "AUC with 95% CI by the DeLong method; P (Wilcoxon) from the two-sided Wilcoxon rank-sum test.",
    "DTL was not detected in GSE111016.",
    "For GPL20301, AUCs are reported for three predefined comparisons; P (Wilcoxon) is shown only for "
    "Sarcopenia+Old vs Young.",
]
write_sheet(ws, "Table 3. Diagnostic performance of the four hub genes across all bulk cohorts.",
            header, rows, notes)

# ---------------- Sheet 4: Table 4 ----------------
ws = wb.create_sheet("Table4_Immune")
ws["A1"] = "Table 4. Immune infiltration analyses (ssGSEA) in the H. pylori line (GSE27411)."
ws["A1"].font = TITLE_FONT

ws.append([])
ws.append(["Part A. ssGSEA immune cell abundance, HP vs Control (GSE27411)."])
ws.cell(row=ws.max_row, column=1).font = Font(name="Arial", size=10, bold=True)
ws.cell(row=ws.max_row, column=1).fill = SEC_FILL
data = read_csv(r"D:\GEO-limma,肌少症1\11-Immune2\output\Immune_groupdiff_HP.csv")
header = ["Cell type", "Control mean", "HP mean", "Diff (HP - Control)", "P", "P adj (BH)"]
rows = data[1:]
start = ws.max_row + 1
ws.append(header)
hr = ws.max_row
for c in range(1, len(header) + 1):
    cell = ws.cell(row=hr, column=c)
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
for r in rows:
    ws.append([fmt(v) for v in r])
    for c in range(1, len(header) + 1):
        ws.cell(row=ws.max_row, column=c).font = BODY_FONT
        ws.cell(row=ws.max_row, column=c).border = BORDER

ws.append([])
ws.append(["Part B. Spearman correlation between hub genes and immune cell abundance, HP line."])
ws.cell(row=ws.max_row, column=1).font = Font(name="Arial", size=10, bold=True)
ws.cell(row=ws.max_row, column=1).fill = SEC_FILL
data = read_csv(r"D:\GEO-limma,肌少症1\11-Immune2\output\Keygene_immune_correlation_HP.csv")
header = ["Gene", "Cell type", "Spearman rho", "P", "P adj (BH)", "Significance"]
rows = [[r[0], r[1], r[2], r[3], r[4], r[5]] for r in data[1:]]
ws.append(header)
hr = ws.max_row
for c in range(1, len(header) + 1):
    cell = ws.cell(row=hr, column=c)
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
for r in rows:
    ws.append([fmt(v) for v in r])
    for c in range(1, len(header) + 1):
        ws.cell(row=ws.max_row, column=c).font = BODY_FONT
        ws.cell(row=ws.max_row, column=c).border = BORDER

ws.append([])
for n in [
    "Part A: P adj by BH across 21 cell types; Macrophages, aDC, DC and Th2 cells were significantly "
    "higher in HP (P adj < 0.05).",
    "Part B: significance codes: * P adj < 0.05, ** P adj < 0.01, *** P adj < 0.001, ns = not significant.",
    "The corresponding sarcopenia-line (GSE1428, n = 22) analyses showed no significant associations and "
    "are provided as supplementary material.",
]:
    ws.append([n])
    ws.cell(row=ws.max_row, column=1).font = NOTE_FONT
for c, w in zip(range(1, 7), [22, 18, 20, 20, 14, 16]):
    ws.column_dimensions[get_column_letter(c)].width = w

wb.save(OUT)
print("saved:", OUT)
for name in wb.sheetnames:
    print(name, "->", wb[name].max_row, "rows")
